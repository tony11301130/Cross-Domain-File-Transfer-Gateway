import os
import io
import zipfile
import shutil
import sys
from PIL import Image, ImageDraw, ImageFont

try:
    import pikepdf
    from reportlab.pdfgen import canvas
    from docx import Document
except ImportError:
    print("Please install required libraries: pip install pikepdf reportlab python-docx")
    sys.exit(1)

# Base directory relative to this script
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SAMPLES_DIR = os.path.join(BASE_DIR, 'samples')

# EICAR Test String
EICAR_STRING = b'X5O!P%@AP[4\\PZX54(P^)7CC)7}$EICAR-STANDARD-ANTIVIRUS-TEST-FILE!$H+H*'

def create_suspicious_pdf(path):
    print(f"Creating realistic suspicious PDF: {path}")
    
    # 1. Create a valid PDF with visible content using ReportLab
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.setFont("Helvetica-Bold", 24)
    c.drawString(100, 700, "CONFIDENTIAL CONTRACT")
    c.setFont("Helvetica", 12)
    c.drawString(100, 650, "This document contains sensitive financial data.")
    c.drawString(100, 630, "Please do not distribute.")
    c.rect(50, 600, 500, 200) # Draw a box
    c.save()
    buf.seek(0)
    
    # 2. Open with pikepdf and inject malice
    pdf = pikepdf.open(buf)
    
    # Inject fake JS with EICAR
    js_code = "app.alert('This is a malicious script!'); // " + EICAR_STRING.decode('ascii')
    
    pdf.Root.Names = pikepdf.Dictionary()
    pdf.Root.Names.JavaScript = pikepdf.Dictionary()
    pdf.Root.Names.JavaScript.BadScript = pikepdf.Dictionary({
        "/S": "/JavaScript",
        "/JS": js_code
    })
    
    # Inject OpenAction (Auto-execute on open)
    pdf.Root.OpenAction = pikepdf.Dictionary({
        "/S": "/JavaScript",
        "/JS": "this.print(true);"
    })
    
    pdf.save(path)

def create_malicious_docx(path):
    print(f"Creating realistic malicious DOCX: {path}")
    
    # 1. Create a real Docx with content
    doc = Document()
    doc.add_heading('Financial Report Q4', 0)
    doc.add_paragraph('This document contains analysis of our quarterly performance.')
    doc.add_paragraph('Please enable macros to view dynamic charts.', style='Quote')
    
    # Create tables/content
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Metric'
    hdr_cells[1].text = 'Value'
    hdr_cells[2].text = 'YoY'
    
    # Save to temp buffer
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    
    # 2. Inject Malware via Zip Manipulation
    # We copy the valid docx structure but INSERT our malicious parts
    with zipfile.ZipFile(path, 'w', zipfile.ZIP_DEFLATED) as zf_out:
        with zipfile.ZipFile(buf, 'r') as zf_in:
            for item in zf_in.infolist():
                # Copy existing files
                zf_out.writestr(item, zf_in.read(item.filename))
        
        # Inject Fake vbaProject.bin with EICAR
        vba_content = b'\xD0\xCF\x11\xE0\xA1\xB1\x1A\xE1' + b'MACRO_HEADER' + EICAR_STRING + b'MACRO_FOOTER'
        zf_out.writestr('word/vbaProject.bin', vba_content)
        
        # Note: To make Word actually "load" the vba, we strictly need to modify [Content_Types].xml 
        # and document.xml.rels. 
        # For CDR demonstration (cleaning), usually just having the file 'vbaProject.bin' is enough 
        # for proper cleaners to trigger.
        # But to be robust, we'll assume the scanner looks for the file existence.
        
        # Inject an OLE Object (Fake entry in zip, might not be referenced in XML but physically exists)
        zf_out.writestr('word/embeddings/oleObject1.bin', b'FAKE_OLE_OBJECT_DATA')

def create_tracked_image(path):
    print(f"Creating realistic tracked Image: {path}")
    
    # Create image with text
    img = Image.new('RGB', (400, 300), color = 'white')
    d = ImageDraw.Draw(img)
    
    # Draw huge red warning
    d.rectangle([10, 10, 390, 290], outline="red", width=5)
    d.text((50, 140), "TOP SECRET DATA", fill="red")
    d.text((50, 160), "DO NOT SHARE", fill="black")
    
    img.save(path, format='JPEG', quality=85) 
    
    # Append EICAR (Steganography)
    with open(path, 'ab') as f:
        f.write(b'__PAYLOAD_START__' + EICAR_STRING + b'__PAYLOAD_END__')

def create_nested_zip(path):
    print(f"Creating nested ZIP attack: {path}")
    
    # Inner RTF
    rtf_content = b'{\\rtf1\\ansi \\b Important Memo \\b0 \\par This file contains hidden threats. {\\object\\objdata ATTACK} }'
    
    inner_zip_io = io.BytesIO()
    with zipfile.ZipFile(inner_zip_io, 'w') as zf:
        zf.writestr('memo.rtf', rtf_content)
        zf.writestr('virus.txt', EICAR_STRING)
    
    with zipfile.ZipFile(path, 'w') as zf:
        zf.writestr('HR_Documents.zip', inner_zip_io.getvalue())
        zf.writestr('ReadMe.txt', b'Please review the attached HR documents.')

if __name__ == "__main__":
    if not os.path.exists(SAMPLES_DIR):
        os.makedirs(SAMPLES_DIR)
        
    create_suspicious_pdf(os.path.join(SAMPLES_DIR, "confidential_contract.pdf"))
    create_malicious_docx(os.path.join(SAMPLES_DIR, "financial_report.docx"))
    create_tracked_image(os.path.join(SAMPLES_DIR, "secret_diagram.jpg"))
    create_nested_zip(os.path.join(SAMPLES_DIR, "hr_bundle.zip"))
    print(f"Realistic samples generated in {SAMPLES_DIR} (with EICAR).")
